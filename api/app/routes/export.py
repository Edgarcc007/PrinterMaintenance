from fastapi import APIRouter, Depends
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session, joinedload
from datetime import date
from io import BytesIO
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from app.database import get_db
from app.models.models import MaintenanceRecord, Printer

router = APIRouter(prefix="/api/export", tags=["Export"])

@router.get("/maintenance")
def export_maintenance(db: Session = Depends(get_db)):
    records = db.query(MaintenanceRecord).options(
        joinedload(MaintenanceRecord.printer).joinedload(Printer.location),
        joinedload(MaintenanceRecord.category)
    ).order_by(MaintenanceRecord.performed_at.desc()).limit(500).all()

    wb = Workbook()
    ws = wb.active
    ws.title = "Maintenance History"

    hdr_font = Font(bold=True, color="FFFFFF", size=11)
    hdr_fill = PatternFill(start_color="2C3E50", end_color="2C3E50", fill_type="solid")
    border = Border(
        left=Side(style="thin"), right=Side(style="thin"),
        top=Side(style="thin"), bottom=Side(style="thin")
    )

    headers = [
        "ID", "Date", "Printer", "Brand", "Serial Number",
        "Location", "Type", "Category", "Performed By",
        "Duration (min)", "Findings", "Actions Taken",
        "Parts Replaced", "Status"
    ]

    for col, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=h)
        cell.font = hdr_font
        cell.fill = hdr_fill
        cell.alignment = Alignment(horizontal="center")
        cell.border = border

    for ri, r in enumerate(records, 2):
        vals = [
            r.id,
            r.performed_at.strftime("%Y-%m-%d %H:%M") if r.performed_at else "",
            f"{r.printer.brand} {r.printer.model}" if r.printer else "",
            r.printer.brand if r.printer else "",
            r.printer.serial_number if r.printer else "",
            r.printer.location.name if r.printer and r.printer.location else "",
            r.type.value if r.type else "",
            r.category.name if r.category else "",
            r.performed_by or "",
            r.duration_minutes or "",
            r.findings or "",
            r.actions_taken or "",
            r.parts_replaced or "",
            r.status or ""
        ]
        for col, val in enumerate(vals, 1):
            cell = ws.cell(row=ri, column=col, value=val)
            cell.border = border

    for col in ws.columns:
        mx = max(len(str(c.value or "")) for c in col)
        ws.column_dimensions[col[0].column_letter].width = min(mx + 3, 40)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)

    return StreamingResponse(
        buf,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f"attachment; filename=maintenance_{date.today()}.xlsx"}
    )
